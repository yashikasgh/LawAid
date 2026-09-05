"""create_template.py — Generates the official Form IF1 FIR template (.docx) with docxtpl tags.

Located at: ai/rag/generation/templates/create_template.py
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def create_fir_template():
    output_dir = Path(__file__).resolve().parent
    output_path = output_dir / "fir_template.docx"

    doc = Document()

    # Page setup - Margins 0.75 inch
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Style definitions
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0, 0, 0)

    # Document Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_if1 = p_header.add_run("FORM I.F.1\nFIRST INFORMATION REPORT\n")
    run_if1.bold = True
    run_if1.font.size = Pt(14)

    run_sub = p_header.add_run("(Under Section 154 Cr.P.C. / Section 173 BNSS)\n")
    run_sub.italic = True
    run_sub.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Field 1: Administrative Identification Table
    t1 = doc.add_table(rows=2, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = t1.rows[0].cells
    hdr_cells[0].text = "District: ______________"
    hdr_cells[1].text = "P.S.: {{jurisdiction_police_station}}"
    hdr_cells[2].text = "Year: ______________"
    hdr_cells[3].text = "FIR No.: ______________"

    row2_cells = t1.rows[1].cells
    row2_cells[0].text = "Date & Time of FIR: ______________"
    row2_cells[1].text = "G.D. Entry No.: ______________"
    row2_cells[2].text = ""
    row2_cells[3].text = ""

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Field 2: Act(s) & Section(s) Title & Loop
    p_f2 = doc.add_paragraph()
    p_f2.add_run("2. Act(s) and Section(s) Involved:").bold = True
    p_f2.paragraph_format.space_after = Pt(2)

    # Dynamic Jinja2 Loop Table for Acts & Sections
    p_loop_start = doc.add_paragraph("{% for item in fir_sections %}")
    p_loop_start.paragraph_format.space_after = Pt(2)

    t_sec = doc.add_table(rows=1, cols=2)
    t_sec.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sec_cells = t_sec.rows[0].cells
    t_sec_cells[0].text = "Act: {{ item.act }}"
    t_sec_cells[1].text = "Section(s): {{ item.sections }}"

    p_loop_end = doc.add_paragraph("{% endfor %}")
    p_loop_end.paragraph_format.space_after = Pt(6)

    # Field 3: Occurrence of Offence & Information Received
    p_f3 = doc.add_paragraph()
    p_f3.add_run("3. Occurrence of Offence and Information Received:").bold = True

    p_f3_a = doc.add_paragraph()
    p_f3_a.add_run("(a) Day of Occurrence: ").bold = True
    p_f3_a.add_run("{{occurrence_day}}\n")
    p_f3_a.add_run("    Date From / To: ").bold = True
    p_f3_a.add_run("{{occurrence_date}}\n")
    p_f3_a.add_run("    Time Period / Time: ").bold = True
    p_f3_a.add_run("{{occurrence_time}}")

    p_f3_b = doc.add_paragraph()
    p_f3_b.add_run("(b) Information Received at Police Station:\n").bold = True
    p_f3_b.add_run("    Date: ").bold = True
    p_f3_b.add_run("{{information_date}}    |    ")
    p_f3_b.add_run("Time: ").bold = True
    p_f3_b.add_run("{{information_time}}")

    p_f3_c = doc.add_paragraph()
    p_f3_c.add_run("(c) General Diary Reference: ").bold = True
    p_f3_c.add_run("Entry No.(s): ______________    |    Date & Time: ______________")

    p_f3_d = doc.add_paragraph()
    p_f3_d.add_run("(d) Type of Information (Written / Oral): ").bold = True
    p_f3_d.add_run("{{information_type}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Field 4: Place of Occurrence
    p_f4 = doc.add_paragraph()
    p_f4.add_run("4. Place of Occurrence:").bold = True

    p_f4_a = doc.add_paragraph()
    p_f4_a.add_run("(a) Distance and Direction from P.S.: ").bold = True
    p_f4_a.add_run("Distance: {{place_distance}}    |    Direction: {{place_direction}}")

    p_f4_b = doc.add_paragraph()
    p_f4_b.add_run("(b) Address / Location: ").bold = True
    p_f4_b.add_run("{{place_address}}")

    p_f4_c = doc.add_paragraph()
    p_f4_c.add_run("(c) Jurisdiction P.S. (If outside local limits): ").bold = True
    p_f4_c.add_run("{{jurisdiction_police_station}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Field 5: Complainant / Informant Details
    p_f5 = doc.add_paragraph()
    p_f5.add_run("5. Complainant / Informant Details:").bold = True

    p_f5_fields = doc.add_paragraph()
    p_f5_fields.add_run("(a) Name: ").bold = True
    p_f5_fields.add_run("{{complainant_name}}\n")
    p_f5_fields.add_run("(b) Father's / Husband's Name: ").bold = True
    p_f5_fields.add_run("{{father_husband_name}}\n")
    p_f5_fields.add_run("(c) Date / Year of Birth: ").bold = True
    p_f5_fields.add_run("{{date_of_birth}}\n")
    p_f5_fields.add_run("(d) Nationality: ").bold = True
    p_f5_fields.add_run("{{nationality}}\n")
    p_f5_fields.add_run("(e) Passport / ID Type: ").bold = True
    p_f5_fields.add_run("{{id_type}}    |    ")
    p_f5_fields.add_run("ID Number: ").bold = True
    p_f5_fields.add_run("{{id_number}}\n")
    p_f5_fields.add_run("(f) Occupation: ").bold = True
    p_f5_fields.add_run("{{occupation}}\n")
    p_f5_fields.add_run("(g) Address: ").bold = True
    p_f5_fields.add_run("{{address}}\n")
    p_f5_fields.add_run("(h) Phone Number: ").bold = True
    p_f5_fields.add_run("{{phone}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Field 6: Details of Known / Suspected / Unknown Accused
    p_f6 = doc.add_paragraph()
    p_f6.add_run("6. Details of Known / Suspected / Unknown Accused with Full Particulars:").bold = True
    p_f6_val = doc.add_paragraph("{{accused_details}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Field 7: Reasons for Delay
    p_f7 = doc.add_paragraph()
    p_f7.add_run("7. Reasons for Delay in Reporting by Complainant / Informant:").bold = True
    p_f7_val = doc.add_paragraph("{{delay_reason}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Field 8: Particulars of Properties Stolen / Involved
    p_f8 = doc.add_paragraph()
    p_f8.add_run("8. Particulars of Properties Stolen / Involved:").bold = True
    p_f8_val = doc.add_paragraph("{{property_details}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Field 9: Total Value of Properties Stolen / Involved
    p_f9 = doc.add_paragraph()
    p_f9.add_run("9. Total Value of Properties Stolen / Involved:").bold = True
    p_f9_val = doc.add_paragraph("{{property_value}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Field 10: Inquest Report / U.D. Case No.
    p_f10 = doc.add_paragraph()
    p_f10.add_run("10. Inquest Report / U.D. Case No., if any: ").bold = True
    p_f10.add_run("______________")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Field 11: First Information Contents
    p_f11 = doc.add_paragraph()
    p_f11.add_run("11. First Information Contents (Attach separate sheet if necessary):").bold = True
    p_f11_val = doc.add_paragraph("{{fir_contents}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Field 12: Action Taken / Investigation Details
    p_f12 = doc.add_paragraph()
    p_f12.add_run("12. Action Taken / Investigation Details:").bold = True
    p_f12_val = doc.add_paragraph("Investigation registered and assigned to Officer in Charge. ______________")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Field 13: Signatures & Administrative Sign-off
    p_f13 = doc.add_paragraph()
    p_f13.add_run("13. Administrative Sign-off and Court Dispatch:").bold = True

    t_sig = doc.add_table(rows=2, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    t_sig_c1 = t_sig.rows[0].cells
    t_sig_c1[0].text = "Signature / Thumb Impression of Complainant:\n\n______________"
    t_sig_c1[1].text = "Signature of Officer in Charge, Police Station:\n\n______________"

    t_sig_c2 = t_sig.rows[1].cells
    t_sig_c2[0].text = "Date & Time of Dispatch to Court:\n______________"
    t_sig_c2[1].text = "Name: ______________\nRank: ______________  No.: ______________"

    # Save to fir_template.docx
    doc.save(str(output_path))
    print(f"Generated official Form IF1 FIR template at: {output_path.resolve()}")
    return output_path


if __name__ == "__main__":
    create_fir_template()
